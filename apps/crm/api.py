"""
API du Suivi patients.

GET    /crm/patients/            liste filtrable (segments, recherche, filtres avances)
GET    /crm/patients/export/     meme liste au format Excel
GET    /crm/patients/search/     recherche rapide d'un patient (nom, telephone, numero)
GET    /crm/passages/            passages recents + decompte par motif
POST   /crm/passages/            enregistrer un passage
DELETE /crm/passages/<id>/       supprimer un passage (auteur sous 24 h, ou administrateur)
"""
import io
from collections import Counter
from datetime import date, timedelta

from django.db.models import Count, OuterRef, Subquery
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from rest_framework import permissions, status
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.core.modules import Modules, user_can
from apps.core.permissions import HasModuleAccess

from . import services
from .models import PatientInteraction


def _base(request):
    """Patients du centre, annotes, filtres selon les parametres de la requete."""
    organization = request.user.organization
    params = request.query_params
    inclure_externes = str(params.get('include_external', '')).lower() in ('1', 'true', 'oui')

    base = services.patients_du_centre(organization, inclure_externes)
    qs = services.avec_activite(base)
    return organization, base, qs


def _liste(request):
    """(organization, base, qs filtre et trie) — partage entre liste et export."""
    organization, base, qs = _base(request)
    params = request.query_params
    montants = services.peut_voir_montants(request.user)

    qs = services.appliquer_segment(qs, params.get('segment') or 'all', organization)
    qs = services.appliquer_filtres(qs, params)
    qs = services.appliquer_ordre(qs, params.get('ordering'), montants)
    return organization, base, qs, montants


class _CrmView(APIView):
    permission_classes = [permissions.IsAuthenticated, HasModuleAccess]
    required_module = Modules.CRM


class CrmPatientListView(_CrmView):
    """Liste des patients avec leur activite, filtrable par segment."""

    def get(self, request):
        try:
            organization, base, qs, montants = _liste(request)
            page = max(int(request.query_params.get('page', 1)), 1)
            taille = min(max(int(request.query_params.get('page_size', 30)), 1), 100)
        except (TypeError, ValueError):
            return Response({'error': 'Paramètre invalide.'}, status=status.HTTP_400_BAD_REQUEST)

        dernier_passage = (PatientInteraction.objects
                           .filter(patient=OuterRef('pk'), kind=PatientInteraction.KIND_PASSAGE)
                           .order_by('-occurred_at'))
        qs = qs.annotate(
            crm_last_passage=Subquery(dernier_passage.values('occurred_at')[:1]),
            crm_last_passage_reason=Subquery(dernier_passage.values('reason')[:1]),
        )

        total = qs.count()
        debut = (page - 1) * taille
        lot = list(qs[debut:debut + taille])

        services_lot = services.services_par_patient([c.id for c in lot])

        # Numeros partages par plusieurs fiches (famille) : le bouton WhatsApp
        # reste utilisable, mais l'ecran le signale pour ne pas envoyer N fois
        # le meme message a la meme personne.
        compte_numeros = Counter()
        for brut in base.exclude(phone='').values_list('phone', flat=True):
            for numero in services.normaliser_telephones(brut)[:1]:
                compte_numeros[numero] += 1

        maintenant = timezone.now()
        libelles_motifs = dict(PatientInteraction.REASON_CHOICES)
        resultats = []
        for c in lot:
            numeros = services.normaliser_telephones(c.phone)
            resultats.append({
                'id': str(c.id),
                'name': c.name,
                'patient_number': c.patient_number,
                'phone': c.phone,
                'whatsapp_url': services.lien_whatsapp(c.phone),
                'phone_shared': compte_numeros[numeros[0]] if numeros else 0,
                'gender': c.gender,
                'age': services.age_en_annees(c.date_of_birth),
                'address': c.address,
                'has_privilege_card': c.has_privilege_card,
                'created_at': c.created_at,
                'visits': c.crm_visits,
                'last_visit': c.crm_last_visit,
                'days_since_last_visit': (maintenant - c.crm_last_visit).days if c.crm_last_visit else None,
                'services': services_lot.get(c.id, []),
                'last_passage': ({
                    'at': c.crm_last_passage,
                    'days': (maintenant - c.crm_last_passage).days,
                    'reason': libelles_motifs.get(c.crm_last_passage_reason, ''),
                } if c.crm_last_passage else None),
                'paid_total': float(c.crm_paid_total or 0) if montants else None,
            })

        compte_segments = services.compter_segments(services.avec_activite(base), organization)

        return Response({
            'count': total,
            'page': page,
            'page_size': taille,
            'results': resultats,
            'segments': [
                {'code': code, 'label': libelle, 'count': compte_segments[code]}
                for code, libelle in services.SEGMENTS
            ],
            'service_choices': [
                {'value': v, 'label': str(l)}
                for v, l in services.Invoice.INVOICE_TYPES if v != 'credit_note'
            ],
            'can_see_amounts': montants,
            'can_export': user_can(request.user, Modules.CRM, 'export'),
        })


class CrmPatientExportView(_CrmView):
    """Export Excel de la liste courante (memes filtres que la liste)."""

    def get(self, request):
        if not user_can(request.user, Modules.CRM, 'export'):
            return Response({'error': "Vous n'avez pas le droit d'exporter."},
                            status=status.HTTP_403_FORBIDDEN)

        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        try:
            organization, base, qs, montants = _liste(request)
        except (TypeError, ValueError):
            return Response({'error': 'Paramètre invalide.'}, status=status.HTTP_400_BAD_REQUEST)

        patients = list(qs[:5000])
        services_lot = services.services_par_patient([c.id for c in patients])

        classeur = Workbook()
        feuille = classeur.active
        feuille.title = 'Suivi patients'
        titres = ['Patient', 'Téléphone', 'Âge', 'Sexe', 'Quartier / adresse',
                  'Visites', 'Dernière visite', 'Services utilisés']
        if montants:
            titres.append('Total payé (FCFA)')
        feuille.append(titres)
        for i in range(1, len(titres) + 1):
            cellule = feuille.cell(row=1, column=i)
            cellule.font = Font(bold=True, color='FFFFFF')
            cellule.fill = PatternFill('solid', fgColor='1F4E79')
            cellule.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        for c in patients:
            ligne = [
                c.name, c.phone or '', services.age_en_annees(c.date_of_birth) or '',
                {'M': 'Homme', 'F': 'Femme'}.get(c.gender, ''), c.address or '',
                c.crm_visits, c.crm_last_visit.date() if c.crm_last_visit else None,
                ', '.join(s['label'] for s in services_lot.get(c.id, [])),
            ]
            if montants:
                ligne.append(int(c.crm_paid_total or 0))
            feuille.append(ligne)

        for i, largeur in enumerate([34, 18, 7, 9, 28, 9, 16, 44, 18][:len(titres)], start=1):
            feuille.column_dimensions[get_column_letter(i)].width = largeur
        for ligne in feuille.iter_rows(min_row=2, min_col=7, max_col=7):
            for cellule in ligne:
                cellule.number_format = 'DD/MM/YYYY'
        feuille.freeze_panes = 'A2'
        feuille.auto_filter.ref = 'A1:%s%s' % (get_column_letter(len(titres)), max(feuille.max_row, 1))

        tampon = io.BytesIO()
        classeur.save(tampon)
        reponse = HttpResponse(
            tampon.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        reponse['Content-Disposition'] = 'attachment; filename="suivi-patients-%s.xlsx"' % date.today()
        return reponse


# ── Recherche rapide d'un patient ───────────────────────────────────────────

class CrmPatientSearchView(_CrmView):
    """Recherche d'un patient pour le bouton « signaler un passage ».

    Propre au Suivi patients : un compte limite a ce module n'a pas le droit
    d'appeler la recherche de la fiche patient (donnees de sante), et il n'en a
    pas besoin — on ne renvoie que de quoi reconnaitre la personne.
    """

    def get(self, request):
        q = (request.query_params.get('q') or '').strip()
        if len(q) < 2:
            return Response({'results': []})
        base = services.patients_du_centre(request.user.organization)
        lot = services.appliquer_filtres(base, {'q': q}).order_by('name')[:10]
        return Response({'results': [{
            'id': str(c.id),
            'name': c.name,
            'patient_number': c.patient_number,
            'phone': c.phone,
            'age': services.age_en_annees(c.date_of_birth),
            'gender': c.gender,
        } for c in lot]})


# ── Passages ────────────────────────────────────────────────────────────────

def _passage_dict(i):
    return {
        'id': str(i.id),
        'patient_id': str(i.patient_id) if i.patient_id else None,
        'name': i.patient.name if i.patient_id else i.person_name,
        'phone': i.patient.phone if i.patient_id else i.person_phone,
        'registered': bool(i.patient_id),
        'reason': i.reason,
        'reason_label': i.get_reason_display(),
        'text': i.text,
        'occurred_at': i.occurred_at,
        'created_by': (i.created_by.get_full_name() or i.created_by.username) if i.created_by_id else '',
    }


class CrmPassageListCreateView(_CrmView):

    def get(self, request):
        params = request.query_params
        try:
            jours = min(max(int(params.get('days', 7)), 1), 365)
            page = max(int(params.get('page', 1)), 1)
            taille = min(max(int(params.get('page_size', 50)), 1), 100)
        except (TypeError, ValueError):
            return Response({'error': 'Paramètre invalide.'}, status=status.HTTP_400_BAD_REQUEST)

        maintenant = timezone.now()
        qs = (PatientInteraction.objects
              .filter(organization=request.user.organization,
                      kind=PatientInteraction.KIND_PASSAGE,
                      occurred_at__gte=maintenant - timedelta(days=jours))
              .select_related('patient', 'created_by'))
        if params.get('patient'):
            qs = qs.filter(patient_id=params['patient'])

        # Le decompte par motif ignore le filtre de motif : il sert justement a choisir.
        compte = dict(qs.values_list('reason').annotate(n=Count('id')))
        if params.get('reason'):
            qs = qs.filter(reason=params['reason'])

        total = qs.count()
        debut = (page - 1) * taille
        aujourdhui = timezone.localdate()
        return Response({
            'count': total,
            'page': page,
            'page_size': taille,
            'results': [_passage_dict(i) for i in qs[debut:debut + taille]],
            'summary': {
                'today': PatientInteraction.objects.filter(
                    organization=request.user.organization, kind=PatientInteraction.KIND_PASSAGE,
                    occurred_at__date=aujourdhui).count(),
                'by_reason': [{'value': v, 'label': l, 'count': compte.get(v, 0)}
                              for v, l in PatientInteraction.REASON_CHOICES],
            },
            'reasons': [{'value': v, 'label': l} for v, l in PatientInteraction.REASON_CHOICES],
        })

    def post(self, request):
        if not user_can(request.user, Modules.CRM, 'create'):
            return Response({'error': "Vous n'avez pas le droit d'enregistrer un passage."},
                            status=status.HTTP_403_FORBIDDEN)
        data = request.data
        organization = request.user.organization

        raison = data.get('reason')
        if raison not in dict(PatientInteraction.REASON_CHOICES):
            return Response({'error': 'Indiquez ce que la personne est venue faire.'},
                            status=status.HTTP_400_BAD_REQUEST)

        patient = None
        if data.get('patient_id'):
            # Un patient de laboratoire partenaire reste selectionnable : il peut
            # passer au centre sans y avoir de dossier.
            patient = services.patients_du_centre(organization, inclure_externes=True).filter(
                id=data['patient_id']).first()
            if patient is None:
                return Response({'error': 'Patient introuvable.'}, status=status.HTTP_404_NOT_FOUND)

        nom = (data.get('person_name') or '').strip()[:150]
        telephone = (data.get('person_phone') or '').strip()[:30]
        if patient is None and not nom:
            return Response({'error': 'Choisissez un patient ou indiquez un nom.'},
                            status=status.HTTP_400_BAD_REQUEST)

        passage = PatientInteraction.objects.create(
            organization=organization, patient=patient,
            person_name='' if patient else nom, person_phone='' if patient else telephone,
            kind=PatientInteraction.KIND_PASSAGE, reason=raison,
            text=(data.get('text') or '').strip()[:500],
            created_by=request.user,
        )
        return Response(_passage_dict(passage), status=status.HTTP_201_CREATED)


class CrmPassageDetailView(_CrmView):

    def delete(self, request, pk):
        passage = get_object_or_404(PatientInteraction, pk=pk, organization=request.user.organization,
                                    kind=PatientInteraction.KIND_PASSAGE)
        auteur_recent = (passage.created_by_id == request.user.id
                         and passage.created_at > timezone.now() - timedelta(hours=24))
        if not (services.peut_voir_montants(request.user) or auteur_recent):
            return Response({'error': "Un passage ne se supprime que dans les 24 h par son auteur, "
                                      "ou par un administrateur."},
                            status=status.HTTP_403_FORBIDDEN)
        passage.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)

