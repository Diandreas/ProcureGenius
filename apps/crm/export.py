"""« Tout télécharger » : toutes les données du Suivi patients dans un seul classeur Excel.

Lecture seule. Les montants n'apparaissent que pour les comptes qui peuvent les voir
(`services.voit_montants`).
"""
import io
from datetime import date, timedelta

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import services
from .models import (CRMCampaign, ContactLog, ContactReason, InvoiceCRMInfo, PatientCRMProfile,
                     PatientInteraction, PatientOrigin)
from .stats import calculer_statistiques

ENTETE = Font(bold=True, color='FFFFFF')
FOND = PatternFill('solid', fgColor='1F4E79')
TEXTE_ETAT = {'came': 'Venu', 'waiting': 'Attendu', 'missed': 'Pas venu', 'none': ''}


def _jour(dt):
    return timezone.localtime(dt).replace(tzinfo=None) if dt else None


def _agent(u):
    return (u.get_full_name() or u.username) if u else ''


def _feuille(classeur, titre, colonnes, lignes, premiere=False):
    f = classeur.active if premiere else classeur.create_sheet()
    f.title = titre[:31]
    f.append([c for c, _ in colonnes])
    for i, (_, largeur) in enumerate(colonnes, start=1):
        cellule = f.cell(row=1, column=i)
        cellule.font, cellule.fill = ENTETE, FOND
        cellule.alignment = Alignment(wrap_text=True, vertical='center')
        f.column_dimensions[get_column_letter(i)].width = largeur
    for ligne in lignes:
        f.append(ligne)
    for colonne in f.iter_cols(min_row=2):
        for cellule in colonne:
            if hasattr(cellule.value, 'hour'):
                cellule.number_format = 'DD/MM/YYYY HH:MM'
            elif isinstance(cellule.value, date):
                cellule.number_format = 'DD/MM/YYYY'
    f.freeze_panes = 'A2'
    if f.max_row > 1:
        f.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(colonnes)), f.max_row)
    return f


def classeur_complet(organization, user):
    montants = services.voit_montants(user)
    maintenant = timezone.now()
    classeur = Workbook()

    # ── Patients ────────────────────────────────────────────────────────────
    base = services.patients_du_centre(organization)
    patients = list(services.avec_activite(base).order_by('name'))
    ids = [p.id for p in patients]
    profils = {p.patient_id: p for p in PatientCRMProfile.objects.filter(patient_id__in=ids)
               .select_related('origin', 'campaign', 'referred_by', 'recorded_by')}
    services_lot = services.services_par_patient(ids)
    nb_messages = {}
    for pid in ContactLog.objects.filter(organization=organization).values_list('patient_id', flat=True):
        nb_messages[pid] = nb_messages.get(pid, 0) + 1
    colonnes = [('Patient', 30), ('N° dossier', 14), ('Téléphone', 16), ('Sexe', 7), ('Âge', 6), ('Date de naissance', 13),
                ('Quartier', 18), ('Adresse', 26), ('Comment il a connu le centre', 22), ('Précision', 20),
                ('Campagne', 18), ('Envoyé par (parrain)', 22), ('Provenance saisie par', 18), ('Ne plus relancer', 9),
                ('Fiche créée le', 16), ('Visites', 8), ('Dernière visite', 16), ('Services utilisés', 40),
                ('Messages reçus', 9)]
    if montants:
        colonnes.append(('Total payé (FCFA)', 14))
    lignes = []
    for c in patients:
        pr = profils.get(c.id)
        ligne = [
            c.name, c.patient_number or '', c.phone or '', {'M': 'Homme', 'F': 'Femme'}.get(c.gender, ''),
            services.age_en_annees(c.date_of_birth), c.date_of_birth,
            services.quartier_de(c, pr) or '', c.address or '',
            (pr.origin.label if pr and pr.origin_id else ('Inconnue' if pr and pr.unknown else '')),
            pr.detail if pr else '', pr.campaign.name if pr and pr.campaign_id else '',
            pr.referred_by.name if pr and pr.referred_by_id else '', _agent(pr.recorded_by) if pr and pr.recorded_by_id else '',
            'oui' if pr and pr.do_not_contact else '', _jour(c.created_at), c.crm_visits, _jour(c.crm_last_visit),
            ', '.join(s['label'] for s in services_lot.get(c.id, [])), nb_messages.get(c.id, 0),
        ]
        if montants:
            ligne.append(int(c.crm_paid_total or 0))
        lignes.append(ligne)
    _feuille(classeur, 'Patients', colonnes, lignes, premiere=True)

    # ── Messages et relances ─────────────────────────────────────────────────
    logs = list(ContactLog.objects.filter(organization=organization)
                .select_related('patient', 'reason', 'campaign', 'created_by').order_by('-contacted_at'))
    etats = services.calculer_venues(logs, maintenant)
    colonnes = [('Date', 16), ('Patient', 28), ('Téléphone', 15), ('Canal', 10), ('Motif', 18), ('Campagne', 16),
                ('Réponse', 22), ('Message envoyé', 60), ('Note', 26), ('Date prévue', 12), ('Envoyé par', 18),
                ('Suite', 10), ('Facture de la venue', 22)]
    if montants:
        colonnes.append(('Montant de la venue (FCFA)', 14))
    lignes = []
    for l in logs:
        etat = etats[l.id]
        facture = etat.get('invoice')
        ligne = [_jour(l.contacted_at), l.patient.name, l.patient.phone or '', l.get_channel_display(),
                 l.reason.label if l.reason_id else '', l.campaign.name if l.campaign_id else '',
                 l.get_outcome_display(), l.message, l.note, l.follow_up_date, _agent(l.created_by),
                 TEXTE_ETAT.get(etat.get('status'), ''), facture.invoice_number if facture else '']
        if montants:
            ligne.append(int(facture.total_amount or 0) if facture else None)
        lignes.append(ligne)
    _feuille(classeur, 'Messages et relances', colonnes, lignes)

    # ── Infos de suivi sur les factures ──────────────────────────────────────
    infos = list(InvoiceCRMInfo.objects.filter(organization=organization)
                 .select_related('invoice', 'invoice__client', 'recorded_by').order_by('-invoice__created_at'))
    colonnes = [('Date facture', 16), ('Facture', 22), ('Patient', 28), ('Venu pour', 26), ('Montée en gamme', 9),
                ('À revoir le', 12), ('Note', 30), ('Saisi par', 18)]
    if montants:
        colonnes[5:5] = [('Montant prévu (FCFA)', 13), ('Total facture (FCFA)', 13)]
    lignes = []
    for i in infos:
        f = i.invoice
        ligne = [_jour(f.created_at), f.invoice_number, f.client.name if f.client_id else '', i.came_for,
                 'oui' if i.upsold else '']
        if montants:
            ligne += [int(i.planned_amount) if i.planned_amount is not None else None, int(f.total_amount or 0)]
        ligne += [i.revisit_date, i.note, _agent(i.recorded_by)]
        lignes.append(ligne)
    _feuille(classeur, 'Infos de suivi factures', colonnes, lignes)

    # ── Passages au centre ───────────────────────────────────────────────────
    motifs_passage = dict(PatientInteraction.REASON_CHOICES)
    pertes = dict(PatientInteraction.LOST_CHOICES)
    lignes = [[_jour(p.occurred_at), p.get_kind_display(), p.patient.name if p.patient_id else p.person_name,
               p.patient.phone if p.patient_id else p.person_phone, motifs_passage.get(p.reason, ''), p.text,
               pertes.get(p.lost_reason, '') if p.lost_reason else 'non', _agent(p.created_by)]
              for p in PatientInteraction.objects.filter(organization=organization)
              .select_related('patient', 'created_by').order_by('-occurred_at')]
    _feuille(classeur, 'Passages et notes', [('Date', 16), ('Type', 9), ('Personne', 28), ('Téléphone', 15),
                                            ('Venu pour', 18), ('Détail', 40), ('Reparti sans acheter', 18),
                                            ('Saisi par', 18)], lignes)

    # ── Campagnes ────────────────────────────────────────────────────────────
    colonnes = [('Campagne', 26), ('Type', 16), ('Zone', 18), ('Début', 12), ('Fin', 12), ('Active', 7),
                ('Patients rattachés', 10), ('Notes', 40)]
    if montants:
        colonnes.insert(5, ('Budget (FCFA)', 12))
    rattaches = {}
    for cid in PatientCRMProfile.objects.filter(organization=organization, campaign__isnull=False).values_list('campaign_id', flat=True):
        rattaches[cid] = rattaches.get(cid, 0) + 1
    lignes = []
    for c in CRMCampaign.objects.filter(organization=organization).order_by('-start_date'):
        ligne = [c.name, c.get_kind_display(), c.zone, c.start_date, c.end_date]
        if montants:
            ligne.append(int(c.budget) if c.budget is not None else None)
        ligne += ['oui' if c.is_active else 'non', rattaches.get(c.id, 0), c.notes]
        lignes.append(ligne)
    _feuille(classeur, 'Campagnes', colonnes, lignes)

    # ── Statistiques des 12 derniers mois ────────────────────────────────────
    fin = timezone.localdate()
    st = calculer_statistiques(organization, fin - timedelta(days=365), fin, montants)
    colonnes = [('Personne', 24), ('Messages envoyés', 10), ('Ont répondu', 10), ("D'accord", 10), ('Venus', 9),
                ('Revenus ensuite', 10), ('% réponse', 9), ("% d'accord", 9), ('% venus', 9), ('% revenus', 9),
                ('% venus / envoyés', 10), ('Passages', 9), ('Passages avec achat', 10), ('Montées en gamme', 10)]
    if montants:
        colonnes.append(('Encaissé après relance (FCFA)', 14))
    lignes = []
    for a in st['team_conversion']:
        ligne = [a['label'], a['sent'], a['answered'], a['agreed'], a['came'], a['retained'], a['answer_rate'],
                 a['agree_rate'], a['came_rate'], a['retained_rate'], a['overall_rate'], a['passages'],
                 a['passages_bought'], a['upsold']]
        if montants:
            ligne.append(a['revenue'])
        lignes.append(ligne)
    _feuille(classeur, 'Conversion équipe (12 mois)', colonnes, lignes)

    colonnes = [('Canal', 24), ('Patients', 9), ('Jamais venus', 9), ('Nouveaux (12 mois)', 10), ('% qui reviennent', 10),
                ('Visites par patient', 10), ('Services les plus utilisés', 40)]
    if montants:
        colonnes += [('Encaissé depuis toujours (FCFA)', 15), ('Par patient (FCFA)', 13), ('Panier moyen (FCFA)', 13),
                     ('Rang rentabilité', 9)]
    lignes = []
    for p in st['provenance']:
        ligne = [p['label'], p['patients'], p['never_came'], p['new_patients'], p['repeat_rate'], p['visits_per_patient'],
                 ', '.join('%s (%d)' % (s['label'], s['count']) for s in p['top_services'])]
        if montants:
            ligne += [p['revenue_lifetime'], p['revenue_per_patient'], p['basket'], p.get('rank')]
        lignes.append(ligne)
    _feuille(classeur, 'Provenance (canaux)', colonnes, lignes)

    # ── Réglages ─────────────────────────────────────────────────────────────
    lignes = [['Provenance', o.label, 'oui' if o.is_active else 'non', ''] for o in PatientOrigin.objects.filter(organization=organization).order_by('position')]
    lignes += [['Motif de relance', m.label, 'oui' if m.is_active else 'non', m.message_template]
               for m in ContactReason.objects.filter(organization=organization).order_by('position')]
    _feuille(classeur, 'Réglages', [('Liste', 18), ('Libellé', 28), ('Actif', 7), ('Message proposé', 70)], lignes)

    tampon = io.BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()
