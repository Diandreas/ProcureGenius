"""Rapport mensuel « patients financiers » : qui a depense quoi, a chaque passage.

Un seul constructeur, utilise par le telechargement (page Suivi patients et page
Patients | Financier) et par l'envoi mensuel par e-mail, pour que le fichier
recu soit identique a celui qu'on telecharge.

Regle financiere (voir memoire project_paiement_mode_source) : « paye » = facture
au statut `paid`. Les factures envoyees mais non reglees apparaissent quand meme
dans le detail, avec leur statut, mais ne sont pas comptees dans le total paye.
"""
import io
from collections import OrderedDict
from datetime import date, timedelta

from django.db.models import Prefetch, Q

from apps.invoicing.models import Invoice, InvoiceItem

from . import services

# Depense minimale visee par patient sur le mois : en dessous, la ligne est surlignee en rouge.
SEUIL_DEPENSE = 20000

STATUTS = {'paid': 'Payée', 'sent': 'Non réglée', 'overdue': 'En retard'}


def mois_precedent(aujourdhui=None):
    """(debut, fin) du mois civil precedent."""
    aujourdhui = aujourdhui or date.today()
    fin = aujourdhui.replace(day=1) - timedelta(days=1)
    return fin.replace(day=1), fin


def bornes_periode(params):
    """(debut, fin) depuis ?start=AAAA-MM-JJ&end=AAAA-MM-JJ, sinon ?month=AAAA-MM, sinon le mois
    precedent. Leve ValueError si invalide."""
    if params.get('start') or params.get('end'):
        debut = date.fromisoformat(str(params.get('start')))
        fin = date.fromisoformat(str(params.get('end')))
        if debut > fin or (fin - debut).days > 400:
            raise ValueError('periode')
        return debut, fin
    if params.get('month'):
        return bornes_du_mois(params['month'])
    return mois_precedent()


def bornes_du_mois(texte):
    """'2026-09' -> (2026-09-01, 2026-09-30). Leve ValueError si invalide."""
    annee, mois = (int(x) for x in texte.split('-'))
    debut = date(annee, mois, 1)
    suivant = date(annee + (mois == 12), (mois % 12) + 1, 1)
    return debut, suivant - timedelta(days=1)


def _libelle_item(item):
    texte = (item.description or '').strip() or 'Article'
    return '%s × %s' % (texte, item.quantity) if item.quantity and item.quantity > 1 else texte


def factures_du_mois(organization, debut, fin, inclure_externes=False):
    patients = services.patients_du_centre(organization, inclure_externes)
    return (Invoice.objects
            .filter(organization=organization, client__in=patients,
                    created_at__date__gte=debut, created_at__date__lte=fin)
            .exclude(status__in=['cancelled', 'draft'])
            .exclude(invoice_type='credit_note')
            .select_related('client')
            .prefetch_related(Prefetch('items', queryset=InvoiceItem.objects.select_related('product__category')
                                       .order_by('created_at')))
            .order_by('created_at'))


def service_de(facture):
    """Service d'une facture (Consultation, Laboratoire, Pharmacie...), comme dans les parcours."""
    from .parcours import SERVICE_PAR_TYPE, _service_ligne
    if facture.invoice_type in SERVICE_PAR_TYPE:
        return SERVICE_PAR_TYPE[facture.invoice_type]
    noms = []
    for i in facture.items.all():
        p = i.product
        cat = p.category.name if p is not None and p.category_id else ''
        s = _service_ligne(cat, p.product_type if p is not None else 'service')
        if s not in noms:
            noms.append(s)
    return ' + '.join(noms) or 'Autres services'


def rangs_des_actions(organization, ids_patients, fin):
    """Pour chaque patient, ses factures depuis toujours (jusqu'a `fin`), dans l'ordre :
    {facture_id: rang (1 = premiere action au centre)}, {patient: [id 1re, id 2e]}."""
    from django.utils import timezone
    from datetime import datetime, time
    limite = timezone.make_aware(datetime.combine(fin, time.max))
    rang, premieres = {}, {}
    for f in (Invoice.objects.filter(organization=organization, client_id__in=ids_patients, created_at__lte=limite)
              .exclude(status__in=['cancelled', 'draft']).exclude(invoice_type='credit_note')
              .order_by('client_id', 'created_at').values('id', 'client_id')):
        liste = premieres.setdefault(f['client_id'], [])
        liste.append(f['id'])
        rang[f['id']] = len(liste)
    return rang, {k: v[:2] for k, v in premieres.items()}


def libelle_rang(n):
    return '1re action' if n == 1 else '%se action' % n


def agreger_par_patient(factures):
    par_patient = OrderedDict()
    for f in factures:
        fiche = par_patient.setdefault(f.client_id, {
            'client': f.client, 'passages': 0, 'paye': 0, 'non_regle': 0, 'articles': OrderedDict(),
        })
        fiche['passages'] += 1
        if f.status == 'paid':
            fiche['paye'] += f.total_amount or 0
        else:
            fiche['non_regle'] += f.total_amount or 0
        for item in f.items.all():
            fiche['articles'][_libelle_item(item)] = True
    return par_patient


TRANCHES_AGE = [(0, 4, '0-4 ans'), (5, 14, '5-14 ans'), (15, 24, '15-24 ans'), (25, 34, '25-34 ans'),
                (35, 44, '35-44 ans'), (45, 59, '45-59 ans'), (60, 200, '60 ans et plus')]
SEXES = {'M': 'Hommes', 'F': 'Femmes'}


def _tranche(age):
    if age is None:
        return 'Âge non renseigné'
    for bas, haut, libelle in TRANCHES_AGE:
        if bas <= age <= haut:
            return libelle
    return 'Âge non renseigné'


def profil_depenses(par_patient, aujourdhui=None):
    """Qui depense : par sexe, par tranche d'age, et les plus gros depensiers.

    Seul le montant paye est compte. Un patient du mois qui n'a rien paye reste
    dans les effectifs (il fait baisser la depense moyenne, c'est voulu).
    """
    total = float(sum(f['paye'] for f in par_patient.values())) or 0

    def bloc(cle, ordre):
        groupes = OrderedDict((lib, {'label': lib, 'patients': 0, 'passages': 0, 'paye': 0.0}) for lib in ordre)
        for fiche in par_patient.values():
            g = groupes.setdefault(cle(fiche['client']), {'label': cle(fiche['client']), 'patients': 0,
                                                            'passages': 0, 'paye': 0.0})
            g['patients'] += 1
            g['passages'] += fiche['passages']
            g['paye'] += float(fiche['paye'])
        sortie = []
        for g in groupes.values():
            if not g['patients']:
                continue
            g['moyenne_par_patient'] = round(g['paye'] / g['patients'])
            g['part'] = round(100 * g['paye'] / total, 1) if total else 0
            sortie.append(g)
        return sortie

    def sexe(c):
        return SEXES.get(c.gender, 'Sexe non renseigné')

    def age(c):
        return _tranche(services.age_en_annees(c.date_of_birth, aujourdhui))

    classes = sorted(par_patient.values(), key=lambda f: -f['paye'])[:10]
    return {
        'by_gender': bloc(sexe, SEXES.values()),
        'by_age': bloc(age, [lib for _, _, lib in TRANCHES_AGE]),
        'top': [{
            'name': f['client'].name, 'patient_number': f['client'].patient_number,
            'age': services.age_en_annees(f['client'].date_of_birth, aujourdhui),
            'gender': SEXES.get(f['client'].gender, ''),
            'passages': f['passages'], 'paye': float(f['paye']),
        } for f in classes],
        'total_paye': total,
        'patients': len(par_patient),
    }


def construire_classeur(organization, debut, fin, inclure_externes=False):
    """Classeur Excel (BytesIO) + totaux {patients, passages, paye}."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from django.utils import timezone
    from .models import PatientCRMProfile

    factures = list(factures_du_mois(organization, debut, fin, inclure_externes))

    par_patient = agreger_par_patient(factures)
    rang, premieres = rangs_des_actions(organization, list(par_patient.keys()), fin)
    ids_premieres = [i for v in premieres.values() for i in v]
    actions = {f.id: f for f in (Invoice.objects.filter(id__in=ids_premieres).select_related('client')
                                 .prefetch_related(Prefetch('items', queryset=InvoiceItem.objects
                                                            .select_related('product__category').order_by('created_at'))))}
    provenances = {}
    for p in PatientCRMProfile.objects.filter(patient_id__in=list(par_patient.keys())).select_related('origin'):
        provenances[p.patient_id] = p.origin.label if p.origin_id else ('Inconnue' if p.unknown else '')
    jour = lambda d: timezone.localtime(d).strftime('%d/%m/%Y')
    heure = lambda d: timezone.localtime(d).strftime('%d/%m/%Y %H:%M')
    dates_periode = {}
    for f in factures:
        dates_periode.setdefault(f.client_id, []).append(f.created_at)

    classeur = Workbook()
    en_tete = Font(bold=True, color='FFFFFF')
    fond = PatternFill('solid', fgColor='1A5276')

    def feuille(ws, titres, lignes, largeurs):
        ws.append(titres)
        for c in ws[1]:
            c.font, c.fill = en_tete, fond
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for ligne in lignes:
            ws.append(ligne)
        for i, l in enumerate(largeurs, 1):
            ws.column_dimensions[get_column_letter(i)].width = l
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical='top', wrap_text=True)
        ws.freeze_panes = 'A2'

    # 1. Resume par patient (ce que demande la direction)
    resume = []
    for fiche in sorted(par_patient.values(), key=lambda x: -x['paye']):
        c = fiche['client']
        age = services.age_en_annees(c.date_of_birth)
        resume.append([c.patient_number or '', c.name, c.phone or '',
                       age if age is not None else '', SEXES.get(c.gender, ''), fiche['passages'],
                       float(fiche['paye']), float(fiche['non_regle']),
                       ' ; '.join(fiche['articles'].keys())])
    # 0. Ce que demande la direction : la 1re et la 2e action de chaque patient venu sur la période
    wa = classeur.active
    wa.title = '1re et 2e action'
    lignes_actions = []
    for pid, fiche in sorted(par_patient.items(), key=lambda kv: min(dates_periode[kv[0]])):
        c = fiche['client']
        ids = premieres.get(pid, [])
        ligne = [c.patient_number or '', c.name, c.phone or '', provenances.get(pid, '')]
        for k in range(2):
            f = actions.get(ids[k]) if k < len(ids) else None
            if f is None:
                ligne += ['', '', '', None, '']
                continue
            ligne += [heure(f.created_at), service_de(f), ' ; '.join(_libelle_item(i) for i in f.items.all()),
                      float(f.total_amount or 0), STATUTS.get(f.status, f.get_status_display())]
        rangs_periode = sorted(rang.get(f.id, 0) for f in factures if f.client_id == pid)
        nouveau = bool(ids) and ids[0] in {f.id for f in factures}
        ligne += ['oui' if nouveau else 'non', ', '.join(libelle_rang(n) for n in rangs_periode),
                  fiche['passages'], float(fiche['paye'])]
        lignes_actions.append(ligne)
    feuille(wa, ['N° patient', 'Patient', 'Téléphone', 'Comment il a connu le centre',
                 '1re action : date', '1re action : service', '1re action : détail', '1re action : montant (FCFA)', '1re action : statut',
                 '2e action : date', '2e action : service', '2e action : détail', '2e action : montant (FCFA)', '2e action : statut',
                 'Nouveau sur la période', 'Actions sur la période', 'Passages sur la période', 'Payé sur la période (FCFA)'],
            lignes_actions, [13, 28, 14, 18, 16, 16, 40, 14, 12, 16, 16, 40, 14, 12, 11, 22, 11, 16])
    vert = PatternFill('solid', fgColor='E2EFDA')
    for row in wa.iter_rows(min_row=2, max_row=1 + len(lignes_actions)):
        if row[14].value == 'oui':
            for c in row[4:9]:
                c.fill = vert
    wa.append([])
    wa.append(['', 'Période du %s au %s. « 1re action » = toute première facture du patient au centre, même avant '
               'la période. En vert : 1re action faite pendant la période (nouveau patient).'
               % (debut.strftime('%d/%m/%Y'), fin.strftime('%d/%m/%Y'))])

    ws = classeur.create_sheet('Par patient')
    for ligne, fiche in zip(resume, sorted(par_patient.values(), key=lambda x: -x['paye'])):
        ds = dates_periode[fiche['client'].id]
        ligne.insert(5, jour(min(ds)))
        ligne.insert(6, jour(max(ds)))
    feuille(ws, ['N° patient', 'Patient', 'Téléphone', 'Âge', 'Sexe', 'Premier passage (période)', 'Dernier passage (période)',
                 'Passages', 'Total payé (FCFA)', 'Non réglé (FCFA)', 'Services / produits'],
            resume, [13, 30, 15, 7, 10, 14, 14, 10, 16, 15, 80])
    rouge = PatternFill('solid', fgColor='FFC7CE')
    for row in ws.iter_rows(min_row=2, max_row=1 + len(resume)):
        if (row[8].value or 0) < SEUIL_DEPENSE:
            for c in row:
                c.fill = rouge
            row[8].font = Font(bold=True, color='9C0006')
    total_paye = float(sum(f['paye'] for f in par_patient.values()))
    ws.append([])
    ws.append(['', 'TOTAL', '', '', '', '', '', len(factures), total_paye,
               float(sum(f['non_regle'] for f in par_patient.values())), ''])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    en_dessous = sum(1 for r in resume if r[8] < SEUIL_DEPENSE)
    ws.append(['', 'En rouge : moins de {:,} FCFA dépensés ({} patient(s) sur {})'.format(
        SEUIL_DEPENSE, en_dessous, len(resume)).replace(',', ' ')])

    # 2. Detail : une ligne par passage (facture)
    detail = []
    for f in factures:
        c = f.client
        n = rang.get(f.id, 0)
        detail.append([
            heure(f.created_at), f.invoice_number, c.patient_number or '', c.name,
            libelle_rang(n) if n else '', 'Première action' if n == 1 else 'Action suivante', service_de(f),
            float(f.total_amount or 0), STATUTS.get(f.status, f.get_status_display()),
            f.get_payment_method_display() if getattr(f, 'payment_method', '') else '',
            ' ; '.join(_libelle_item(i) for i in f.items.all()),
        ])
    profil = profil_depenses(par_patient)
    wp = classeur.create_sheet('Profil des dépenses')
    wp.append(['Qui dépense ?'])
    wp['A1'].font = Font(bold=True, size=13)
    for titre_bloc, cle in (('Par sexe', 'by_gender'), ('Par tranche d\'âge', 'by_age')):
        wp.append([])
        wp.append([titre_bloc, 'Patients', 'Passages', 'Total payé (FCFA)', 'Part du total (%)',
                   'Dépense moyenne par patient (FCFA)'])
        for c in wp[wp.max_row]:
            c.font, c.fill = en_tete, fond
        for g in profil[cle]:
            wp.append([g['label'], g['patients'], g['passages'], g['paye'], g['part'], g['moyenne_par_patient']])
    wp.append([])
    wp.append(['10 plus gros dépensiers', 'Âge', 'Sexe', 'Total payé (FCFA)', 'Passages'])
    for c in wp[wp.max_row]:
        c.font, c.fill = en_tete, fond
    for x in profil['top']:
        wp.append([x['name'], x['age'] if x['age'] is not None else '', x['gender'], x['paye'], x['passages']])
    for i, l in enumerate([34, 12, 12, 18, 16, 30], 1):
        wp.column_dimensions[get_column_letter(i)].width = l

    ws2 = classeur.create_sheet('Détail des passages')
    feuille(ws2, ['Date et heure', 'N° facture', 'N° patient', 'Patient', 'N° de l\'action', 'Type d\'action', 'Service',
                  'Montant (FCFA)', 'Statut', 'Mode de paiement', 'Services / produits'],
            detail, [16, 18, 13, 30, 11, 15, 16, 15, 13, 16, 80])

    tampon = io.BytesIO()
    classeur.save(tampon)
    tampon.seek(0)
    return tampon, {'patients': len(par_patient), 'passages': len(factures), 'paye': total_paye,
                    'sous_seuil': en_dessous}


def nom_fichier(organization, debut, fin=None):
    if fin is None:
        return 'patients_financiers_%s_%s.xlsx' % (organization.name.replace(' ', '_'), debut.strftime('%Y-%m'))
    return 'patients_financiers_%s_%s_au_%s.xlsx' % (organization.name.replace(' ', '_'),
                                                     debut.strftime('%Y-%m-%d'), fin.strftime('%Y-%m-%d'))


def envoyer_rapport(organization, debut, fin, destinataires, inclure_externes=False):
    """Envoie le fichier du mois par e-mail avec le SMTP de l'organisation.

    Leve RuntimeError (message lisible) si la configuration d'envoi manque.
    """
    from django.core.mail import EmailMessage

    from apps.core.email_utils import get_organization_email_backend
    from apps.accounts.models import EmailConfiguration

    config = EmailConfiguration.objects.filter(organization=organization).first()
    connexion = get_organization_email_backend(organization)
    if not config or not connexion:
        raise RuntimeError("Aucune configuration d'envoi d'e-mails (SMTP) pour cette organisation.")
    if not destinataires:
        raise RuntimeError('Aucun destinataire.')

    tampon, totaux = construire_classeur(organization, debut, fin, inclure_externes)
    mois = debut.strftime('%m/%Y')
    corps = (
        'Bonjour,\n\nCi-joint le fichier des patients financiers de %s : pour chaque patient, '
        'le montant dépensé à chaque passage et les services ou produits achetés.\n\n'
        '%s patient(s), %s passage(s), %s FCFA payés.\n\n'
        'Envoyé automatiquement par ProcureGenius.'
    ) % (mois, totaux['patients'], totaux['passages'], '{:,.0f}'.format(totaux['paye']).replace(',', ' '))
    message = EmailMessage(
        subject='[%s] Patients financiers — %s' % (organization.name, mois),
        body=corps,
        from_email='%s <%s>' % (config.default_from_name, config.default_from_email),
        to=destinataires, connection=connexion,
    )
    message.attach(nom_fichier(organization, debut), tampon.read(),
                   'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    message.send()
    return totaux
